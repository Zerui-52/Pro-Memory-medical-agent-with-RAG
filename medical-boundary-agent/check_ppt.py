from pptx import Presentation
from lxml import etree

prs = Presentation(r'C:\Users\24991\Desktop\Modern Redesign (Same Content).pptx')

# Check spTree structure of slide 2
slide = prs.slides[1]
spTree = slide.shapes._spTree

# Find all grpSp (group shape) elements
ns = {
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
}

groups = spTree.findall('.//p:grpSp', ns)
print(f"Slide 2 has {len(groups)} group shapes")

for i, grp in enumerate(groups):
    cNvPr = grp.find('./p:nvGrpSpPr/p:cNvPr', ns)
    if cNvPr is not None:
        gid = cNvPr.get("id")
        name = cNvPr.get("name")
        xfrm = grp.find('./p:grpSpPr/a:xfrm', ns)
        off = xfrm.find('a:off', ns) if xfrm is not None else None
        ext = xfrm.find('a:ext', ns) if xfrm is not None else None
        pos = f"({off.get('x')},{off.get('y')})" if off is not None else "unknown"
        size = f"({ext.get('cx')},{ext.get('cy')})" if ext is not None else "unknown"
        print(f"  Group {i+1}: id={gid}, name={name}, pos={pos}, size={size}")

# Check slide 5 groups  
slide5 = prs.slides[4]
spTree5 = slide5.shapes._spTree
groups5 = spTree5.findall('.//p:grpSp', ns)
print(f"\nSlide 5 has {len(groups5)} group shapes")
for i, grp in enumerate(groups5):
    cNvPr = grp.find('./p:nvGrpSpPr/p:cNvPr', ns)
    if cNvPr is not None:
        gid = cNvPr.get("id")
        name = cNvPr.get("name")
        xfrm = grp.find('./p:grpSpPr/a:xfrm', ns)
        off = xfrm.find('a:off', ns) if xfrm is not None else None
        ext = xfrm.find('a:ext', ns) if xfrm is not None else None
        pos = f"({off.get('x')},{off.get('y')})" if off is not None else "unknown"
        print(f"  Group {i+1}: id={gid}, name={name}, pos={pos}")
